import docx
import os

files = [
    "Primer Seismoisolation.docx",
    "Метод дискретизации уравнений колебаний.docx",
]

for filename in files:
    if not os.path.exists(filename):
        print(f"File not found: {filename}")
        continue
    print(f"\n{'='*80}")
    print(f"FILE: {filename}")
    print('='*80)
    doc = docx.Document(filename)
    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            print(text)
    # Also read tables
    for i, table in enumerate(doc.tables):
        print(f"\n[TABLE {i+1}]")
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells)
            if row_text.strip(" |"):
                print(row_text)
